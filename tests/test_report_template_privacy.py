from pathlib import Path
import unittest

from openpyxl import load_workbook


TEMPLATE_FILE = Path(__file__).resolve().parents[1] / "当日数据统计_模板.xlsx"


class ReportTemplatePrivacyTests(unittest.TestCase):
    def test_non_auckland_template_uses_current_board_capacities(self):
        workbook = load_workbook(TEMPLATE_FILE, read_only=True, data_only=False)
        sheet = workbook["非奥克兰"]

        self.assertEqual(
            [sheet["H2"].value, sheet["I2"].value],
            [280, "4L预测板数"],
        )
        self.assertEqual(
            [sheet["H13"].value, sheet["I13"].value],
            [350, "5L预测板数"],
        )

    def test_template_contains_no_runtime_waybill_rows(self):
        workbook = load_workbook(TEMPLATE_FILE, read_only=True, data_only=False)
        sheet = workbook["数据源1-预测"]

        populated_runtime_cells = [
            cell.value
            for row in sheet.iter_rows(min_row=2)
            for cell in row
            if cell.value not in (None, "")
        ]

        self.assertEqual(populated_runtime_cells, [])

    def test_unused_legacy_data_tabs_are_empty(self):
        workbook = load_workbook(TEMPLATE_FILE, read_only=True, data_only=False)

        for sheet_name in ["Sheet3", "Sheet2", "Sheet1", "数据源2-route code"]:
            with self.subTest(sheet=sheet_name):
                populated = [
                    cell.value
                    for row in workbook[sheet_name].iter_rows()
                    for cell in row
                    if cell.value not in (None, "")
                ]
                self.assertEqual(populated, [])


if __name__ == "__main__":
    unittest.main()
